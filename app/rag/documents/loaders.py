"""Document loaders (TXT/MD, PDF, DOCX) plus optional ``<file>.meta.json`` sidecar metadata."""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.core.exceptions import DocumentProcessingError

SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx"}


@dataclass
class RawDocument:
    path: Path
    #: One entry per page for paginated formats (PDF); a single entry otherwise.
    pages: list[str]
    paginated: bool = False
    metadata: dict[str, Any] = field(default_factory=dict)


def _load_text(path: Path) -> tuple[list[str], bool]:
    try:
        return [path.read_text(encoding="utf-8")], False
    except UnicodeDecodeError:
        return [path.read_text(encoding="latin-1")], False


def _load_pdf(path: Path) -> tuple[list[str], bool]:
    from pypdf import PdfReader

    pages = [page.extract_text() or "" for page in PdfReader(str(path)).pages]
    if not any(p.strip() for p in pages):
        raise DocumentProcessingError(f"{path.name}: no extractable text (scanned PDF? OCR required)")
    return pages, True


def _load_docx(path: Path) -> tuple[list[str], bool]:
    import docx

    document = docx.Document(str(path))
    return ["\n".join(p.text for p in document.paragraphs)], False


_LOADERS = {".txt": _load_text, ".md": _load_text, ".pdf": _load_pdf, ".docx": _load_docx}


def load_sidecar(path: Path) -> dict[str, Any]:
    sidecar = path.with_name(path.name + ".meta.json")
    if not sidecar.exists():
        return {}
    try:
        return json.loads(sidecar.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise DocumentProcessingError(f"invalid metadata sidecar {sidecar.name}: {exc}") from exc


def load_document(path: Path) -> RawDocument:
    suffix = path.suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        raise DocumentProcessingError(f"unsupported document type {suffix!r} ({path.name})")
    try:
        pages, paginated = loader(path)
    except DocumentProcessingError:
        raise
    except Exception as exc:
        raise DocumentProcessingError(f"failed to read {path.name}: {type(exc).__name__}: {exc}") from exc
    return RawDocument(path=path, pages=pages, paginated=paginated, metadata=load_sidecar(path))


def discover(source: Path) -> list[Path]:
    if source.is_file():
        return [source]
    if not source.is_dir():
        raise DocumentProcessingError(f"source {source} does not exist")
    return sorted(p for p in source.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES)


def slugify(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-") or "document"
